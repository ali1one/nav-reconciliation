"""
nav_break_computation.py
========================
Author  : Youssef Ait Ali
Purpose : Step 2 of the NAV reconciliation. For every fund present in all
          three sources (output of Step 1), compare the Internal NAV with the
          Fund Administrator and the Custodian, measure the differences in
          basis points, apply tolerance thresholds and flag breaks with a
          probable root cause.

Usage
-----
    python nav_break_computation.py
    python nav_break_computation.py --data-dir samples/breaks_demo
    python nav_break_computation.py --warn-bps 1 --break-bps 5 --shares-bps 1

Logic
-----
    Reference source = INTERNAL. Two comparisons per fund:
        INTERNAL vs ADMIN
        INTERNAL vs CUSTODIAN

    For each comparison:
        navps_diff_bps   = (NAVPS_counterparty / NAVPS_internal - 1) * 10 000
        shares_diff_bps  = (Units_counterparty / Units_internal - 1) * 10 000
        total_nav_diff   = TotalNAV_counterparty - TotalNAV_internal  (fund ccy)

    Status of a comparison:
        MATCHED  |navps| <= warn  and |shares| <= shares tolerance
        WARNING  warn < |navps| <= break
        BREAK    |navps| > break, or |shares| > shares tolerance,
                 or NAV dates / static data disagree

    Probable root cause:
        PRICING           NAV per share differs, unit count agrees
        CAPITAL_ACTIVITY  unit count differs (subscriptions / redemptions)
        PRICING+CAPITAL   both differ
        NAV_DATE          sources are not valued on the same date
        STATIC_DATA       currency or share class disagree

    Triangulation (who is probably wrong):
        both counterparties break against INTERNAL and agree with each other
            -> INTERNAL is suspect
        only one counterparty breaks -> that counterparty is suspect

Output
------
    output/step2_breaks_YYYY-MM-DD.xlsx
        Sheet "Summary"      one row per fund, worst status, suspect source
        Sheet "Break Detail" one row per fund x comparison, all metrics
        Sheet "Exceptions"   only WARNING / BREAK rows, with action notes

Dependencies : pandas, openpyxl  (re-uses the loaders of nav_coverage_check.py)
"""

from __future__ import annotations

import argparse
import logging
import os
from dataclasses import dataclass
from datetime import date

import pandas as pd

import nav_coverage_check as step1

log = logging.getLogger(__name__)

COUNTERPARTIES = ["ADMIN", "CUSTODIAN"]
STATUS_RANK = {"MATCHED": 0, "WARNING": 1, "BREAK": 2}


# ─────────────────────────────────────────────────────────────────────────────
# CONFIG — TOLERANCES
# ─────────────────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class Tolerance:
    navps_warn_bps: float = 1.0    # above this: WARNING
    navps_break_bps: float = 5.0   # above this: BREAK
    shares_bps: float = 1.0        # unit count difference above this: BREAK


# Tighter or looser thresholds for specific funds (ISIN -> Tolerance).
# Example: money market funds usually get a much tighter tolerance.
TOLERANCE_OVERRIDES: dict[str, Tolerance] = {
    "LU0006789012": Tolerance(navps_warn_bps=0.5, navps_break_bps=2.0, shares_bps=1.0),
}


# ─────────────────────────────────────────────────────────────────────────────
# STEP 1 — LOAD (re-use Step 1 loaders) AND KEEP FULLY COVERED FUNDS
# ─────────────────────────────────────────────────────────────────────────────

def load_sources(data_dir: str | None = None) -> dict[str, pd.DataFrame]:
    """Load the three sources with the Step 1 loaders (optionally from another folder)."""
    if data_dir:
        step1.DATA_DIR = os.path.abspath(data_dir)
    return {
        "INTERNAL": step1.load_internal(),
        "ADMIN": step1.load_admin(),
        "CUSTODIAN": step1.load_custodian(),
    }


def build_wide(sources: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, list[str]]:
    """
    One row per ISIN present in all three sources, with suffixed columns
    (navps_ADMIN, shares_CUSTODIAN, ...). Returns (wide, excluded_isins).
    """
    cols = ["isin", "fund_name", "share_class", "currency",
            "nav_date", "total_nav", "shares", "navps"]

    frames = []
    for src, df in sources.items():
        dup = df["isin"][df["isin"].duplicated()]
        if not dup.empty:
            raise ValueError(f"[{src}] duplicated ISIN(s): {sorted(dup.unique())}")
        frames.append(df[cols].set_index("isin").add_suffix(f"_{src}"))

    all_isins = set().union(*(set(df["isin"]) for df in sources.values()))
    wide = pd.concat(frames, axis=1, join="inner").reset_index()
    excluded = sorted(all_isins - set(wide["isin"]))
    return wide.sort_values("isin").reset_index(drop=True), excluded


# ─────────────────────────────────────────────────────────────────────────────
# STEP 2 — COMPUTE BREAKS
# ─────────────────────────────────────────────────────────────────────────────

def _bps(counterparty: pd.Series, reference: pd.Series) -> pd.Series:
    return (counterparty / reference - 1.0) * 10_000


def _tolerance_for(isin: str, default: Tolerance) -> Tolerance:
    return TOLERANCE_OVERRIDES.get(isin, default)


def _assess(row: pd.Series, tol: Tolerance) -> tuple[str, str]:
    """Return (status, root_cause) for one fund x counterparty comparison."""
    if not row["same_nav_date"]:
        return "BREAK", "NAV_DATE"
    if not row["same_static"]:
        return "BREAK", "STATIC_DATA"

    navps_abs = abs(row["navps_diff_bps"])
    shares_out = abs(row["shares_diff_bps"]) > tol.shares_bps

    if navps_abs > tol.navps_break_bps:
        navps_status = "BREAK"
    elif navps_abs > tol.navps_warn_bps:
        navps_status = "WARNING"
    else:
        navps_status = "MATCHED"

    status = "BREAK" if shares_out else navps_status
    if status == "MATCHED":
        return status, ""

    pricing = navps_status != "MATCHED"
    if pricing and shares_out:
        return status, "PRICING+CAPITAL"
    return status, "CAPITAL_ACTIVITY" if shares_out else "PRICING"


def compute_breaks(wide: pd.DataFrame, tol: Tolerance = Tolerance()) -> pd.DataFrame:
    """One row per ISIN x counterparty with all difference metrics and a status."""
    rows = []
    for cp in COUNTERPARTIES:
        d = pd.DataFrame({
            "isin": wide["isin"],
            "fund_name": wide["fund_name_INTERNAL"],
            "share_class": wide["share_class_INTERNAL"],
            "currency": wide["currency_INTERNAL"],
            "comparison": f"INTERNAL vs {cp}",
            "counterparty": cp,
            "nav_date_internal": wide["nav_date_INTERNAL"],
            "nav_date_counterparty": wide[f"nav_date_{cp}"],
            "navps_internal": wide["navps_INTERNAL"],
            "navps_counterparty": wide[f"navps_{cp}"],
            "navps_diff_bps": _bps(wide[f"navps_{cp}"], wide["navps_INTERNAL"]),
            "shares_internal": wide["shares_INTERNAL"],
            "shares_counterparty": wide[f"shares_{cp}"],
            "shares_diff": wide[f"shares_{cp}"] - wide["shares_INTERNAL"],
            "shares_diff_bps": _bps(wide[f"shares_{cp}"], wide["shares_INTERNAL"]),
            "total_nav_internal": wide["total_nav_INTERNAL"],
            "total_nav_counterparty": wide[f"total_nav_{cp}"],
            "total_nav_diff": wide[f"total_nav_{cp}"] - wide["total_nav_INTERNAL"],
            "total_nav_diff_bps": _bps(wide[f"total_nav_{cp}"], wide["total_nav_INTERNAL"]),
            "same_nav_date": wide["nav_date_INTERNAL"] == wide[f"nav_date_{cp}"],
            "same_static": (wide["currency_INTERNAL"] == wide[f"currency_{cp}"])
                           & (wide["share_class_INTERNAL"] == wide[f"share_class_{cp}"]),
        })
        rows.append(d)

    breaks = pd.concat(rows, ignore_index=True)
    assessed = breaks.apply(
        lambda r: _assess(r, _tolerance_for(r["isin"], tol)), axis=1, result_type="expand"
    )
    breaks["status"], breaks["root_cause"] = assessed[0], assessed[1]
    return breaks.sort_values(["isin", "counterparty"]).reset_index(drop=True)


# ─────────────────────────────────────────────────────────────────────────────
# STEP 3 — TRIANGULATE & SUMMARISE PER FUND
# ─────────────────────────────────────────────────────────────────────────────

ACTION = {
    "INTERNAL": "Both counterparties agree with each other but not with Internal. "
                "Review internal pricing / capital activity bookings.",
    "ADMIN": "Only the Administrator disagrees. Request the Administrator's NAV pack "
             "and price / unit details.",
    "CUSTODIAN": "Only the Custodian disagrees. Check the Custodian's prices and "
                 "settled subscriptions / redemptions.",
    "UNCLEAR": "All three sources disagree. Escalate and investigate position by position.",
}


def _counterparties_agree(wide_row: pd.Series, tol: Tolerance) -> bool:
    navps = abs(_bps(wide_row["navps_CUSTODIAN"], wide_row["navps_ADMIN"]))
    shares = abs(_bps(wide_row["shares_CUSTODIAN"], wide_row["shares_ADMIN"]))
    return navps <= tol.navps_warn_bps and shares <= tol.shares_bps


def summarise(breaks: pd.DataFrame, wide: pd.DataFrame,
              tol: Tolerance = Tolerance()) -> pd.DataFrame:
    """One row per fund: worst status, root causes and the probable faulty source."""
    out = []
    wide_idx = wide.set_index("isin")
    for isin, g in breaks.groupby("isin", sort=True):
        worst = max(g["status"], key=STATUS_RANK.get)
        flagged = g[g["status"] != "MATCHED"]
        suspect, action = "", ""

        if not flagged.empty:
            if len(flagged) == 2:
                agree = _counterparties_agree(wide_idx.loc[isin], _tolerance_for(isin, tol))
                suspect = "INTERNAL" if agree else "UNCLEAR"
            else:
                suspect = flagged["counterparty"].iloc[0]
            action = ACTION[suspect]

        g_adm = g.set_index("counterparty")
        out.append({
            "isin": isin,
            "fund_name": g["fund_name"].iloc[0],
            "share_class": g["share_class"].iloc[0],
            "currency": g["currency"].iloc[0],
            "navps_diff_bps_ADMIN": g_adm.loc["ADMIN", "navps_diff_bps"],
            "navps_diff_bps_CUSTODIAN": g_adm.loc["CUSTODIAN", "navps_diff_bps"],
            "status_ADMIN": g_adm.loc["ADMIN", "status"],
            "status_CUSTODIAN": g_adm.loc["CUSTODIAN", "status"],
            "overall_status": worst,
            "root_cause": ", ".join(sorted(set(flagged["root_cause"]))),
            "suspect_source": suspect,
            "action": action,
        })
    return pd.DataFrame(out)


# ─────────────────────────────────────────────────────────────────────────────
# STEP 4 — CONSOLE REPORT
# ─────────────────────────────────────────────────────────────────────────────

def print_report(summary: pd.DataFrame, excluded: list[str],
                 nav_date: date, tol: Tolerance) -> None:
    W = 98
    icon = {"MATCHED": "✓", "WARNING": "!", "BREAK": "✗"}
    print(f"\n{'═' * W}")
    print("  FUND NAV RECONCILIATION — STEP 2: BREAK COMPUTATION")
    print(f"  NAV Date : {nav_date}  |  Tolerances : NAVPS warn {tol.navps_warn_bps} bp, "
          f"break {tol.navps_break_bps} bp, units {tol.shares_bps} bp"
          f"  |  Overrides : {len(TOLERANCE_OVERRIDES)}")
    print(f"{'═' * W}\n")

    print(f"  {'ISIN':<14} {'Fund Name':<28} {'CCY':<4} "
          f"{'ADM bp':>8} {'CUST bp':>8}  {'Status':<9} {'Suspect':<10} Cause")
    print("  " + "─" * 94)
    for _, r in summary.iterrows():
        print(f"  {r['isin']:<14} {str(r['fund_name'])[:27]:<28} {r['currency']:<4} "
              f"{r['navps_diff_bps_ADMIN']:>8.2f} {r['navps_diff_bps_CUSTODIAN']:>8.2f}  "
              f"{icon[r['overall_status']]} {r['overall_status']:<7} "
              f"{r['suspect_source']:<10} {r['root_cause']}")

    counts = summary["overall_status"].value_counts()
    print("\nSUMMARY")
    print("─" * W)
    print(f"  Funds reconciled            : {len(summary)}")
    for s in STATUS_RANK:
        print(f"  {icon[s]}  {s:<24}: {counts.get(s, 0)}")
    if excluded:
        print(f"  Excluded (incomplete, see Step 1): {len(excluded)}  →  {', '.join(excluded)}")
    print(f"{'═' * W}\n")


# ─────────────────────────────────────────────────────────────────────────────
# STEP 5 — EXCEL EXPORT
# ─────────────────────────────────────────────────────────────────────────────

def export_excel(summary: pd.DataFrame, breaks: pd.DataFrame,
                 nav_date: date, output_dir: str) -> str:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fills = {
        "header": PatternFill("solid", fgColor="1F3864"),
        "MATCHED": PatternFill("solid", fgColor="C6EFCE"),
        "WARNING": PatternFill("solid", fgColor="FFEB9C"),
        "BREAK": PatternFill("solid", fgColor="FFC7CE"),
    }
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    detail = breaks.drop(columns=["counterparty"])
    exceptions = summary[summary["overall_status"] != "MATCHED"]

    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"step2_breaks_{nav_date}.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        detail.to_excel(writer, sheet_name="Break Detail", index=False)
        exceptions.to_excel(writer, sheet_name="Exceptions", index=False)

    wb = load_workbook(path)
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        for c in ws[1]:
            c.fill, c.font, c.border = fills["header"], Font(bold=True, color="FFFFFF", size=9), border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.border = Font(size=9), border
                h = headers[c.column - 1]
                if isinstance(c.value, float):
                    c.number_format = "0.00" if "bps" in h else "#,##0.00"
                if h in ("status", "overall_status", "status_ADMIN", "status_CUSTODIAN") \
                        and c.value in fills:
                    c.fill = fills[c.value]
                if h == "action":
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, 50)
        ws.freeze_panes = "B2"
    wb.save(path)
    return path


# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def run(data_dir: str | None, output_dir: str, tol: Tolerance) -> str:
    sources = load_sources(data_dir)
    wide, excluded = build_wide(sources)
    if wide.empty:
        raise SystemExit("No fund is present in all three sources — run Step 1 first.")

    nav_date = wide["nav_date_INTERNAL"].iloc[0]
    breaks = compute_breaks(wide, tol)
    summary = summarise(breaks, wide, tol)

    print_report(summary, excluded, nav_date, tol)
    path = export_excel(summary, breaks, nav_date, output_dir)
    log.info(f"Output → {path}")
    return path


def main() -> None:
    p = argparse.ArgumentParser(description="NAV reconciliation — Step 2: break computation")
    p.add_argument("--data-dir", help="folder containing internal/ admin/ custodian/ "
                                      "(default: data/)")
    p.add_argument("--output-dir", default=step1.OUTPUT_DIR)
    p.add_argument("--warn-bps", type=float, default=Tolerance.navps_warn_bps)
    p.add_argument("--break-bps", type=float, default=Tolerance.navps_break_bps)
    p.add_argument("--shares-bps", type=float, default=Tolerance.shares_bps)
    a = p.parse_args()
    if a.warn_bps > a.break_bps:
        p.error("--warn-bps must be <= --break-bps")
    run(a.data_dir, a.output_dir, Tolerance(a.warn_bps, a.break_bps, a.shares_bps))


if __name__ == "__main__":
    main()
